"""
把 Daily.xlsx 中每个 Member 的数据追加到文件夹下对应人员“最新一周”的文件中。
- 选择规则：同名文件中 W4 > W3 > W2 > W1（若有多个月份，取最新月份的最大周）
- 以 Invoice No 去重，已存在则不追加
- 增量第 1~18 列 -> 原文件第 1~18 列；增量第 19 列 -> 原文件第 25 列
- 保留原文件格式：新行复制上一数据行的样式，并确保 Status 下拉框覆盖新行
"""
import os
import re
from copy import copy

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

# ===================== 配置 =====================
BASE_DIR = "./data"
DAILY_FILE = os.path.join(BASE_DIR, "Daily.xlsx")
SHEET_NAME = None            # 目标 sheet 名，None 表示第一个 sheet
HEADER_ROW = 1               # 表头所在行

INVOICE_HEADER = "Invoice No"
MEMBER_HEADER = "Member"
STATUS_HEADER = "Status"

COMMON_COLS = 18             # 增量前 18 列 -> 原文件前 18 列
EXTRA_SRC_COL = 19           # 增量第 19 列
EXTRA_DST_COL = 25           # -> 原文件第 25 列
TOTAL_COLS = 30              # 原文件总列数
# ===============================================

FILE_PATTERN = re.compile(r"^(\d{6})-W([1-4])-(.+)\.xlsx$", re.IGNORECASE)


def norm(v):
    """统一 Invoice No / 名字的比较格式（去空格；12345.0 -> '12345'）"""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def find_col(ws, header_name):
    for cell in ws[HEADER_ROW]:
        if norm(cell.value).lower() == header_name.lower():
            return cell.column
    raise ValueError(f"{ws.title} 中找不到表头 '{header_name}'")


def find_latest_files(folder):
    """返回 {member(小写): 路径}，每个人取 (月份, 周) 最大的文件"""
    best = {}
    for fn in os.listdir(folder):
        if fn.startswith("~$"):          # 跳过 Excel 临时锁文件
            continue
        m = FILE_PATTERN.match(fn)
        if not m:
            continue
        month, week, name = m.group(1), int(m.group(2)), m.group(3).strip()
        key = name.lower()
        if key not in best or (month, week) > best[key][0]:
            best[key] = ((month, week), os.path.join(folder, fn))
    return {k: v[1] for k, v in best.items()}


def read_daily(path):
    """按 Member 分组读取增量数据，每行取前 19 列的值"""
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    header = [norm(c) for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    member_idx = [h.lower() for h in header].index(MEMBER_HEADER.lower())
    grouped = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row is None or all(v is None for v in row):
            continue
        row = list(row) + [None] * (EXTRA_SRC_COL - len(row))
        member = norm(row[member_idx])
        if member:
            grouped.setdefault(member.lower(), []).append(row[:EXTRA_SRC_COL])
    wb.close()
    return grouped


def last_data_row(ws, key_col):
    """以 Invoice 列判断最后一个有数据的行（避免被空白格式行误导）"""
    for r in range(ws.max_row, HEADER_ROW, -1):
        if ws.cell(r, key_col).value not in (None, ""):
            return r
    return HEADER_ROW


def ensure_validation(ws, status_col, new_last_row):
    """如果 Status 下拉框的范围没有覆盖到新行，则把范围扩展到新行"""
    col_letter = get_column_letter(status_col)
    for dv in ws.data_validations.dataValidation:
        for rng in list(dv.sqref.ranges):
            if rng.min_col <= status_col <= rng.max_col and rng.max_row < new_last_row:
                dv.sqref.add(f"{col_letter}{rng.max_row + 1}:{col_letter}{new_last_row}")
                return
            if rng.min_col <= status_col <= rng.max_col:
                return  # 已覆盖


def append_to_file(path, rows):
    wb = load_workbook(path)   # 不能用 data_only，否则会丢公式
    ws = wb[SHEET_NAME] if SHEET_NAME else wb.worksheets[0]

    inv_col = find_col(ws, INVOICE_HEADER)
    status_col = find_col(ws, STATUS_HEADER)
    inv_src_idx = inv_col - 1  # 前 18 列顺序一致，所以下标相同

    last = last_data_row(ws, inv_col)
    existing = {norm(ws.cell(r, inv_col).value) for r in range(HEADER_ROW + 1, last + 1)}
    existing.discard("")

    style_row = last if last > HEADER_ROW else None
    added = skipped = 0
    for row in rows:
        inv = norm(row[inv_src_idx])
        if not inv or inv in existing:
            skipped += 1
            continue
        last += 1
        values = {c: row[c - 1] for c in range(1, COMMON_COLS + 1)}
        values[EXTRA_DST_COL] = row[EXTRA_SRC_COL - 1]
        for c in range(1, TOTAL_COLS + 1):
            cell = ws.cell(last, c)
            if style_row:                         # 复制上一数据行的格式
                src = ws.cell(style_row, c)
                if src.has_style:
                    cell._style = copy(src._style)
            if c in values:
                v = values[c]
                # Status 下拉框为 1/2/3，统一成整数
                if c == status_col and isinstance(v, (str, float)) and norm(v).isdigit():
                    v = int(norm(v))
                cell.value = v
        existing.add(inv)
        added += 1

    if added:
        if style_row and ws.row_dimensions[style_row].height:
            for r in range(style_row + 1, last + 1):
                ws.row_dimensions[r].height = ws.row_dimensions[style_row].height
        ensure_validation(ws, status_col, last)
        wb.save(path)
    wb.close()
    return added, skipped


def main():
    files = find_latest_files(BASE_DIR)
    daily = read_daily(DAILY_FILE)

    for member, rows in daily.items():
        path = files.get(member)
        if not path:
            print(f"[跳过] {member}: 文件夹下没有对应文件")
            continue
        added, skipped = append_to_file(path, rows)
        print(f"[完成] {os.path.basename(path)}: 追加 {added} 行, 跳过重复 {skipped} 行")

    for member in files:
        if member not in daily:
            print(f"[跳过] {member}: 增量文件中无数据")


if __name__ == "__main__":
    main()
