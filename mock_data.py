"""生成模拟数据：目标文件夹下的周文件 + Daily.xlsx 增量文件"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

BASE_DIR = "./data"                                  # 周文件所在文件夹
DAILY_FILE = os.path.join(BASE_DIR, "Daily.xlsx")    # 增量文件

# 原文件 30 列：前 18 列与增量文件前 18 列相同，第 25 列对应增量文件第 19 列
COMMON_HEADERS = ["Invoice No", "Member", "Date", "Customer", "Status",
                  "Amount", "Currency", "Region", "Product", "Qty",
                  "Unit Price", "Tax", "Discount", "Channel", "Owner",
                  "Due Date", "Payment Term", "Remark"]            # 18 列
TARGET_HEADERS = COMMON_HEADERS + [f"Extra{i}" for i in range(19, 25)] \
                 + ["Comment"] + [f"Extra{i}" for i in range(26, 31)]  # 30 列
DAILY_HEADERS = COMMON_HEADERS + ["Comment"]                          # 19 列

thin = Side(style="thin", color="999999")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def make_row(inv, member, status, comment=None, extra=False):
    row = [inv, member, "2026-09-01", "Cust-" + inv[-2:], status,
           1000.5, "USD", "APAC", "P1", 3, 333.5, 0.06, 0, "Online",
           "Owner", "2026-10-01", "Net30", "ok"]
    if extra:  # 原文件行：30 列
        row += [None] * 6 + [comment] + [None] * 5
    else:      # 增量文件行：19 列
        row += [comment]
    return row


def make_week_file(path, member, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(TARGET_HEADERS)
    for c in ws[1]:
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", start_color="1F4E78")
        c.alignment = Alignment(horizontal="center")
        c.border = BORDER
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name="Arial")
            c.border = BORDER
        row[5].number_format = "#,##0.00"   # Amount
        row[11].number_format = "0%"        # Tax
    # Status 下拉框（第 5 列 E），预设到 E2:E200
    dv = DataValidation(type="list", formula1='"1,2,3"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("E2:E200")
    ws.freeze_panes = "A2"
    wb.save(path)


def main():
    os.makedirs(BASE_DIR, exist_ok=True)
    # Anna 有 W1~W4，Lancy 只有 W1~W3，Bob 只有 W1
    plan = {"Anna": ["W1", "W2", "W3", "W4"], "Lancy": ["W1", "W2", "W3"], "Bob": ["W1"]}
    for member, weeks in plan.items():
        for i, w in enumerate(weeks, 1):
            rows = [make_row(f"INV-{member[:2].upper()}-{i}{k}", member, 1,
                             "old", extra=True) for k in range(2)]
            make_week_file(os.path.join(BASE_DIR, f"202609-{w}-{member}.xlsx"), member, rows)

    # 增量文件：Anna 含一条已存在的 Invoice（应跳过），Tom 没有对应文件，Bob 无增量数据
    wb = Workbook()
    ws = wb.active
    ws.append(DAILY_HEADERS)
    ws.append(make_row("INV-AN-40", "Anna", 2, "dup, should skip"))   # W4 已有
    ws.append(make_row("INV-AN-NEW1", "Anna", 3, "new anna 1"))
    ws.append(make_row("INV-AN-NEW1", "Anna", 3, "dup in daily"))      # 增量内自身重复
    ws.append(make_row("INV-LA-NEW1", "Lancy", 2, "new lancy 1"))
    ws.append(make_row("INV-LA-NEW2", "Lancy", 1, "new lancy 2"))
    ws.append(make_row("INV-TO-NEW1", "Tom", 1, "no file for Tom"))
    wb.save(DAILY_FILE)
    print("模拟数据已生成:", os.path.abspath(BASE_DIR))


if __name__ == "__main__":
    main()
